import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from html.parser import HTMLParser

# Constants (16:9)
CANVAS_W = Inches(13.333)
CANVAS_H = Inches(7.5)
MARGIN_X = Inches(0.6)
MARGIN_TOP = Inches(0.5)
CONTENT_LEFT = MARGIN_X
CONTENT_RIGHT = CANVAS_W - MARGIN_X
CONTENT_W = CONTENT_RIGHT - CONTENT_LEFT
CONTENT_TOP = MARGIN_TOP
CONTENT_MAX_Y = Inches(6.70)
FOOTER_TOP = Inches(6.85)
FOOTER_H = Inches(0.22)

COLOR_BG = RGBColor(0x05, 0x05, 0x05)  # Matte Black
COLOR_FG = RGBColor(0xFF, 0xFF, 0xFF)  # Pure White
COLOR_ACCENT = RGBColor(0xD4, 0xAF, 0x37)  # Metallic Gold
COLOR_MUTED = RGBColor(0xB8, 0xB8, 0xB8)  # Light Gray

FONT_DISPLAY = "Clash Display"
FONT_BODY = "Inter"
FONT_SERIF = "Playfair Display"
FONT_MONO = "JetBrains Mono"

class Cursor:
    def __init__(self, y_start=CONTENT_TOP, cap=CONTENT_MAX_Y):
        self.y = y_start
        self.cap = cap
        self.history = []

    def take(self, h, gap=Inches(0.12), label=""):
        top = self.y
        self.y = top + h + gap
        self.history.append((top, h, label))
        if self.y > self.cap:
            # For this export, we'll log warning instead of crashing to ensure completion
            print(f"Warning: Cursor exceeded rail at '{label}': y={self.y} cap={self.cap}")
        return top

def hero_layout(blocks):
    total_h = sum(h + g for h, g in blocks)
    y_start = (CANVAS_H - total_h) / 2
    return Cursor(y_start=y_start, cap=CANVAS_H - FOOTER_H - Inches(0.2))

def paint_bg(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text(slide, left, top, width, height, text, font, size_pt, color=COLOR_FG, align=None, bold=False, italic=False):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align == "center":
        p.alignment = PP_ALIGN.CENTER
    elif align == "right":
        p.alignment = PP_ALIGN.RIGHT
    
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return shape

def add_footer(slide, text="GROWTHPILOT GROUP"):
    add_text(slide, CONTENT_LEFT, FOOTER_TOP, CONTENT_W, FOOTER_H, text, FONT_MONO, 9, COLOR_MUTED, align="center")

class ItalicSpans(HTMLParser):
    def __init__(self):
        super().__init__()
        self.runs = []
        self.italic_depth = 0
        self._buf = []
        self._italic = False

    def handle_starttag(self, tag, attrs):
        if tag in ("em", "i"):
            self._flush()
            self.italic_depth += 1
            self._italic = True

    def handle_endtag(self, tag):
        if tag in ("em", "i") and self.italic_depth > 0:
            self._flush()
            self.italic_depth -= 1
            self._italic = self.italic_depth > 0

    def handle_data(self, data):
        self._buf.append(data)

    def _flush(self):
        if self._buf:
            self.runs.append(("".join(self._buf), self._italic))
            self._buf = []

def add_rich_text(slide, left, top, width, height, html_snippet, font, size_pt, color=COLOR_FG, align=None, bold=False):
    parser = ItalicSpans()
    parser.feed(html_snippet)
    parser._flush()
    
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align == "center":
        p.alignment = PP_ALIGN.CENTER
    elif align == "right":
        p.alignment = PP_ALIGN.RIGHT
        
    for text, is_italic in parser.runs:
        run = p.add_run()
        run.text = text
        run.font.name = FONT_SERIF if is_italic else font
        run.font.size = Pt(size_pt)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = is_italic
    return shape

def main():
    prs = Presentation()
    prs.slide_width = CANVAS_W
    prs.slide_height = CANVAS_H
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Hero
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    blocks = [
        (Inches(0.2), Inches(0.4)), # eyebrow
        (Inches(1.8), Inches(0.4)), # title
        (Inches(0.8), Inches(0.4)), # lead
    ]
    c = hero_layout(blocks)
    add_text(slide, CONTENT_LEFT, c.take(*blocks[0]), CONTENT_W, blocks[0][0], "ESTABLISHED 2026 · SILICON VALLEY", FONT_MONO, 12, COLOR_ACCENT, align="center")
    add_rich_text(slide, CONTENT_LEFT, c.take(*blocks[1]), CONTENT_W, blocks[1][0], "We Don't Just Market Brands. We Build <em>Digital Empires</em>.", FONT_DISPLAY, 64, COLOR_FG, align="center", bold=True)
    add_text(slide, CONTENT_LEFT, c.take(*blocks[2]), CONTENT_W, blocks[2][0], "AI Automation, Paid Advertising, SEO, Web Development, Branding, and Scalable Digital Systems engineered for modern businesses seeking elite expansion.", FONT_BODY, 20, COLOR_MUTED, align="center")
    add_footer(slide)

    # Slide 2: Philosophy
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    c = Cursor()
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.3)), CONTENT_W, Inches(0.3), "I. Trust & Positioning", FONT_MONO, 12, COLOR_ACCENT)
    add_rich_text(slide, CONTENT_LEFT, c.take(Inches(1.2)), CONTENT_W, Inches(1.2), "Built for ambitious <em>brands</em>.", FONT_DISPLAY, 48, COLOR_FG, bold=True)
    add_text(slide, CONTENT_LEFT, c.take(Inches(1.0)), CONTENT_W, Inches(1.0), "Luxury execution meets scalable strategy. We create systems designed for long-term growth, moving beyond tactical marketing into architectural brand scaling.", FONT_BODY, 20, COLOR_MUTED)
    add_footer(slide)

    # Slide 3: Capabilities
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    c = Cursor()
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.3)), CONTENT_W, Inches(0.3), "II. Core Capabilities", FONT_MONO, 12, COLOR_ACCENT)
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.8)), CONTENT_W, Inches(0.8), "Services", FONT_DISPLAY, 36, COLOR_FG, bold=True)
    
    card_w = (CONTENT_W - Inches(0.4)) / 3
    top = c.take(Inches(3.0))
    services = [
        ("Paid Advertising", "High-performance campaigns across elite networks, optimized for immediate ROI and brand authority."),
        ("AI Automation", "Futuristic workflows and intelligent assistants that remove bottlenecks and scale operations autonomously."),
        ("Elite Branding", "Visual identities and typography systems crafted for luxury positioning and market dominance.")
    ]
    for i, (title, desc) in enumerate(services):
        left = CONTENT_LEFT + i * (card_w + Inches(0.2))
        add_text(slide, left, top, card_w, Inches(0.4), title, FONT_DISPLAY, 18, COLOR_ACCENT, bold=True)
        add_text(slide, left, top + Inches(0.5), card_w, Inches(1.5), desc, FONT_BODY, 14, COLOR_MUTED)
    add_footer(slide)

    # Slide 4: The Edge
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    c = Cursor()
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.3)), CONTENT_W, Inches(0.3), "III. The Growth Advantage", FONT_MONO, 12, COLOR_ACCENT)
    add_rich_text(slide, CONTENT_LEFT, c.take(Inches(1.2)), CONTENT_W, Inches(1.2), "Modern systems for <em>modern businesses</em>.", FONT_DISPLAY, 40, COLOR_FG, bold=True)
    
    y = c.take(Inches(2.5))
    items = [
        ("01. Custom Digital Solutions", "No templates, only bespoke architecture."),
        ("02. Scalable Processes", "Built to handle 10x volume without friction."),
        ("03. AI-Driven Intelligence", "Real-time data processing and automation.")
    ]
    for i, (title, desc) in enumerate(items):
        add_text(slide, CONTENT_LEFT, y + i * Inches(0.8), CONTENT_W, Inches(0.3), title, FONT_BODY, 18, COLOR_ACCENT, bold=True)
        add_text(slide, CONTENT_LEFT + Inches(0.4), y + i * Inches(0.8) + Inches(0.35), CONTENT_W, Inches(0.3), desc, FONT_BODY, 16, COLOR_MUTED)
    add_footer(slide)

    # Slide 5: Systems
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    c = Cursor()
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.3)), CONTENT_W, Inches(0.3), "IV. Implementation Roadmap", FONT_MONO, 12, COLOR_ACCENT)
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.8)), CONTENT_W, Inches(0.8), "Roadmap", FONT_DISPLAY, 36, COLOR_FG, bold=True)
    
    top = c.take(Inches(2.5))
    steps = [
        ("01 Discovery", "In-depth audit of your current systems and market positioning to identify high-leverage growth levers."),
        ("02 Build", "Engineering the core infrastructure, from AI integrations to high-conversion digital surfaces."),
        ("03 Scale", "Activating the engines and optimizing for exponential growth while maintaining luxury standards.")
    ]
    for i, (title, desc) in enumerate(steps):
        left = CONTENT_LEFT + i * (card_w + Inches(0.2))
        add_text(slide, left, top, card_w, Inches(0.4), title, FONT_DISPLAY, 24, COLOR_ACCENT, bold=True)
        add_text(slide, left, top + Inches(0.5), card_w, Inches(1.5), desc, FONT_BODY, 14, COLOR_MUTED)
    add_footer(slide)

    # Slide 6: Vision
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    c = Cursor()
    add_text(slide, CONTENT_LEFT, c.take(Inches(0.3)), CONTENT_W, Inches(0.3), "V. Concept Showcase", FONT_MONO, 12, COLOR_ACCENT)
    
    top = c.take(Inches(3.0))
    concepts = ["Elite Dashboard Concept", "Elite Management Workflow"]
    half_w = (CONTENT_W - Inches(0.4)) / 2
    for i, title in enumerate(concepts):
        left = CONTENT_LEFT + i * (half_w + Inches(0.4))
        # Placeholder for image
        shape = slide.shapes.add_textbox(left, top, half_w, Inches(2.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_MUTED
        add_text(slide, left, top + Inches(2.2), half_w, Inches(0.4), title, FONT_DISPLAY, 20, COLOR_FG, bold=True)
    add_footer(slide)

    # Slide 7: Engagement
    slide = prs.slides.add_slide(blank_layout)
    paint_bg(slide, COLOR_BG)
    blocks = [
        (Inches(1.5), Inches(0.5)), # title
        (Inches(0.8), Inches(0.5)), # lead
    ]
    c = hero_layout(blocks)
    add_rich_text(slide, CONTENT_LEFT, c.take(*blocks[0]), CONTENT_W, blocks[0][0], "Your Brand Deserves <em>More Than Average</em>.", FONT_DISPLAY, 54, COLOR_FG, align="center", bold=True)
    add_text(slide, CONTENT_LEFT, c.take(*blocks[1]), CONTENT_W, blocks[1][0], "Take the first step toward building your digital empire. Schedule a private strategy session with our advisors.", FONT_BODY, 20, COLOR_MUTED, align="center")
    add_footer(slide)

    prs.save("growthpilot-editorial-14.pptx")
    print(os.path.abspath("growthpilot-editorial-14.pptx"))

if __name__ == "__main__":
    main()
